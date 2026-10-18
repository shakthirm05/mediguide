"""Import MediGuide's fictional directory and recurring schedule templates.

Run from the project root after applying 002_hospital_directory_columns.sql:
    python import_hospital_directory.py MediGuide_Final_Hospital_Dataset_Expanded.xlsx

This imports only Departments, Doctors, and Doctor_Schedules. It deliberately
DOES NOT import Appointment_Slot_Fixtures as live availability or create bookings.
"""
from __future__ import annotations

import argparse
import os
from datetime import date, datetime, time
from pathlib import Path
from uuid import uuid4

import openpyxl
import psycopg
from dotenv import load_dotenv
from psycopg.rows import dict_row


WEEKDAYS = {
    "MON": 1,
    "TUE": 2,
    "WED": 3,
    "THU": 4,
    "FRI": 5,
    "SAT": 6,
    "SUN": 7,
}


def as_date(value, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid date in {field}: {value!r}") from exc


def as_time(value, field: str) -> time:
    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)
    if isinstance(value, time):
        return value.replace(tzinfo=None)
    try:
        return time.fromisoformat(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid time in {field}: {value!r}") from exc


def read_sheet(workbook, sheet_name: str) -> list[dict]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Required worksheet is missing: {sheet_name}")
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = next(rows)
    result = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None for value in values):
            continue
        record = dict(zip(headers, values))
        record["__row_number"] = row_number
        result.append(record)
    return result


def required_text(row: dict, key: str, sheet: str) -> str:
    value = row.get(key)
    if value is None or not str(value).strip():
        raise ValueError(
            f"{sheet} row {row['__row_number']}: {key} is required"
        )
    return str(value).strip()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", help="Path to expanded MediGuide XLSX workbook")
    args = parser.parse_args()

    load_dotenv(Path.cwd() / ".env")
    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL is missing from the project-root .env")

    workbook_path = Path(args.workbook).expanduser().resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        departments = read_sheet(workbook, "Departments")
        doctors = read_sheet(workbook, "Doctors")
        schedules = read_sheet(workbook, "Doctor_Schedules")
    finally:
        workbook.close()

    # Validate the workbook's internal relationships before touching PostgreSQL.
    department_codes = {
        required_text(row, "department_id", "Departments") for row in departments
    }
    doctor_codes = {
        required_text(row, "doctor_id", "Doctors") for row in doctors
    }
    schedule_codes = [
        required_text(row, "schedule_id", "Doctor_Schedules") for row in schedules
    ]
    if len(department_codes) != len(departments):
        raise ValueError("Duplicate department_id values in workbook")
    if len(doctor_codes) != len(doctors):
        raise ValueError("Duplicate doctor_id values in workbook")
    if len(set(schedule_codes)) != len(schedule_codes):
        raise ValueError("Duplicate schedule_id values in workbook")

    for row in doctors:
        dept_code = required_text(row, "department_id", "Doctors")
        if dept_code not in department_codes:
            raise ValueError(
                f"Doctor {row.get('doctor_id')} references missing department {dept_code}"
            )
    for row in schedules:
        doctor_code = required_text(row, "doctor_id", "Doctor_Schedules")
        dept_code = required_text(row, "department_id", "Doctor_Schedules")
        if doctor_code not in doctor_codes:
            raise ValueError(
                f"Schedule {row.get('schedule_id')} references missing doctor {doctor_code}"
            )
        if dept_code not in department_codes:
            raise ValueError(
                f"Schedule {row.get('schedule_id')} references missing department {dept_code}"
            )
        weekday = required_text(row, "weekday", "Doctor_Schedules").upper()[:3]
        if weekday not in WEEKDAYS:
            raise ValueError(
                f"Schedule {row.get('schedule_id')} has invalid weekday {weekday!r}"
            )
        start = as_time(row.get("start_time"), "start_time")
        end = as_time(row.get("end_time"), "end_time")
        if end <= start:
            raise ValueError(
                f"Schedule {row.get('schedule_id')} must end after it starts"
            )
        slot_minutes = int(row.get("slot_minutes") or 0)
        if not 5 <= slot_minutes <= 240:
            raise ValueError(
                f"Schedule {row.get('schedule_id')} has invalid slot_minutes"
            )

    department_uuid_by_code: dict[str, str] = {}
    doctor_uuid_by_code: dict[str, str] = {}

    with psycopg.connect(database_url, row_factory=dict_row) as conn:
        with conn.cursor() as cur:
            # Upsert department rows by stable workbook source ID.
            for row in departments:
                code = required_text(row, "department_id", "Departments")
                name = required_text(row, "department_name", "Departments")
                location = required_text(row, "location_code", "Departments")
                floor = str(row.get("floor") or "").strip() or None
                access_class = str(row.get("access_class") or "").strip() or None
                service_scope = str(row.get("service_scope") or "").strip() or None
                cur.execute(
                    """
                    INSERT INTO departments (
                        department_id, source_code, name, location_code, floor_label,
                        access_class, service_scope, timezone_name, is_active
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, 'Asia/Kolkata', TRUE)
                    ON CONFLICT (source_code) WHERE source_code IS NOT NULL
                    DO UPDATE SET
                        name = EXCLUDED.name,
                        location_code = EXCLUDED.location_code,
                        floor_label = EXCLUDED.floor_label,
                        access_class = EXCLUDED.access_class,
                        service_scope = EXCLUDED.service_scope,
                        is_active = TRUE
                    RETURNING department_id
                    """,
                    (str(uuid4()), code, name, location, floor,
                     access_class, service_scope),
                )
                department_uuid_by_code[code] = str(cur.fetchone()["department_id"])

            # Also load the mapping for any existing source-coded department rows.
            cur.execute(
                "SELECT source_code, department_id FROM departments WHERE source_code IS NOT NULL"
            )
            department_uuid_by_code.update({
                row["source_code"]: str(row["department_id"])
                for row in cur.fetchall()
            })

            for row in doctors:
                code = required_text(row, "doctor_id", "Doctors")
                name = " ".join(required_text(row, "doctor_name", "Doctors").split())
                dept_code = required_text(row, "department_id", "Doctors")
                specialty = str(row.get("specialty") or "").strip() or None
                location = str(row.get("location_code") or "").strip() or None
                cur.execute(
                    """
                    INSERT INTO doctors (
                        doctor_id, source_code, department_id, display_name,
                        specialty, location_code, is_active
                    ) VALUES (%s, %s, %s, %s, %s, %s, TRUE)
                    ON CONFLICT (source_code) WHERE source_code IS NOT NULL
                    DO UPDATE SET
                        department_id = EXCLUDED.department_id,
                        display_name = EXCLUDED.display_name,
                        specialty = EXCLUDED.specialty,
                        location_code = EXCLUDED.location_code,
                        is_active = TRUE
                    RETURNING doctor_id
                    """,
                    (str(uuid4()), code, department_uuid_by_code[dept_code],
                     name, specialty, location),
                )
                doctor_uuid_by_code[code] = str(cur.fetchone()["doctor_id"])

            cur.execute(
                "SELECT source_code, doctor_id FROM doctors WHERE source_code IS NOT NULL"
            )
            doctor_uuid_by_code.update({
                row["source_code"]: str(row["doctor_id"])
                for row in cur.fetchall()
            })

            for row in schedules:
                source_id = required_text(row, "schedule_id", "Doctor_Schedules")
                doctor_code = required_text(row, "doctor_id", "Doctor_Schedules")
                department_code = required_text(row, "department_id", "Doctor_Schedules")
                weekday = required_text(row, "weekday", "Doctor_Schedules").upper()[:3]
                start = as_time(row.get("start_time"), "start_time")
                end = as_time(row.get("end_time"), "end_time")
                slot_minutes = int(row["slot_minutes"])
                valid_from = as_date(row.get("effective_from"), "effective_from")
                valid_until_value = row.get("effective_to")
                valid_until = (
                    as_date(valid_until_value, "effective_to")
                    if valid_until_value not in (None, "") else None
                )
                location = str(row.get("location_code") or "").strip() or None
                booking_mode = str(row.get("booking_mode") or "").strip() or None
                availability_status = str(row.get("availability_status") or "").strip() or None
                notes = str(row.get("notes") or "").strip() or None
                cur.execute(
                    """
                    INSERT INTO doctor_schedules (
                        schedule_id, source_schedule_id, doctor_id, department_id,
                        location_code, weekday, start_time, end_time,
                        slot_duration_minutes, valid_from, valid_until,
                        booking_mode, availability_status, notes, is_active
                    ) VALUES (
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, TRUE
                    )
                    ON CONFLICT (source_schedule_id) WHERE source_schedule_id IS NOT NULL
                    DO UPDATE SET
                        doctor_id = EXCLUDED.doctor_id,
                        department_id = EXCLUDED.department_id,
                        location_code = EXCLUDED.location_code,
                        weekday = EXCLUDED.weekday,
                        start_time = EXCLUDED.start_time,
                        end_time = EXCLUDED.end_time,
                        slot_duration_minutes = EXCLUDED.slot_duration_minutes,
                        valid_from = EXCLUDED.valid_from,
                        valid_until = EXCLUDED.valid_until,
                        booking_mode = EXCLUDED.booking_mode,
                        availability_status = EXCLUDED.availability_status,
                        notes = EXCLUDED.notes,
                        is_active = TRUE
                    """,
                    (str(uuid4()), source_id, doctor_uuid_by_code[doctor_code],
                     department_uuid_by_code[department_code], location,
                     WEEKDAYS[weekday], start, end, slot_minutes, valid_from,
                     valid_until, booking_mode, availability_status, notes),
                )

            # Final consistency checks run before the transaction commits.
            cur.execute("SELECT COUNT(*) AS n FROM departments WHERE source_code IS NOT NULL")
            department_count = cur.fetchone()["n"]
            cur.execute("SELECT COUNT(*) AS n FROM doctors WHERE source_code IS NOT NULL")
            doctor_count = cur.fetchone()["n"]
            cur.execute("SELECT COUNT(*) AS n FROM doctor_schedules WHERE source_schedule_id IS NOT NULL")
            schedule_count = cur.fetchone()["n"]

            expected = (len(departments), len(doctors), len(schedules))
            actual = (department_count, doctor_count, schedule_count)
            if actual != expected:
                raise RuntimeError(
                    f"Imported source-record counts mismatch. Expected {expected}, got {actual}."
                )

    print("MediGuide directory import completed.")
    print(f"Departments: {len(departments)}")
    print(f"Doctors: {len(doctors)}")
    print(f"Recurring schedule templates: {len(schedules)}")
    print("Appointment fixtures imported as live bookings: 0 (intentionally)")
    print("Note: schedules are fictional recurring templates, not live availability.")


if __name__ == "__main__":
    main()
